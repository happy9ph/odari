"""Grades, exports (Excel/PDF) and report cards."""
import random
from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import Float, ForeignKey, String, UniqueConstraint, func, or_, select
from sqlalchemy.orm import Mapped, Session, joinedload, mapped_column

from main import Base, Student, get_db, guard

SUBJECTS = ["Mathematics", "French", "English", "Science", "History", "Geography"]
TERMS = ["Term 1", "Term 2", "Term 3"]


class Grade(Base):
    __tablename__ = "grades"
    __table_args__ = (UniqueConstraint("student_id", "subject", "term"),)
    id: Mapped[int] = mapped_column(primary_key=True)
    student_id: Mapped[int] = mapped_column(ForeignKey("students.id", ondelete="CASCADE"))
    subject: Mapped[str] = mapped_column(String(50))
    term: Mapped[str] = mapped_column(String(20))
    score: Mapped[float] = mapped_column(Float)


def seed_extras(db: Session):
    if not db.scalar(select(func.count(Grade.id))):
        rnd = random.Random(7)
        for s in db.scalars(select(Student)):
            for sub in SUBJECTS:
                for t in TERMS[:2]:
                    db.add(Grade(student_id=s.id, subject=sub, term=t, score=round(rnd.uniform(7, 19) * 4) / 4))
    db.commit()


router = APIRouter(prefix="/api", dependencies=[Depends(guard)])


class ScoreIn(BaseModel):
    student_id: int
    score: float | None = Field(None, ge=0, le=20)


class GradesIn(BaseModel):
    subject: str
    term: str
    scores: list[ScoreIn]


@router.get("/meta")
def meta():
    return dict(subjects=SUBJECTS, terms=TERMS)


@router.get("/grades")
def get_grades(class_id: int, subject: str, term: str, db: Session = Depends(get_db)):
    studs = db.scalars(select(Student).where(Student.class_id == class_id).order_by(Student.full_name)).all()
    sc = {g.student_id: g.score for g in db.scalars(select(Grade).where(
        Grade.subject == subject, Grade.term == term, Grade.student_id.in_([s.id for s in studs] or [0])))}
    return [dict(student_id=s.id, student_number=s.student_number, full_name=s.full_name, score=sc.get(s.id)) for s in studs]


@router.put("/grades")
def save_grades(d: GradesIn, db: Session = Depends(get_db)):
    if d.subject not in SUBJECTS or d.term not in TERMS:
        raise HTTPException(400, "Unknown subject or term")
    ex = {g.student_id: g for g in db.scalars(select(Grade).where(Grade.subject == d.subject, Grade.term == d.term))}
    for s in d.scores:
        g = ex.get(s.student_id)
        if s.score is None:
            if g:
                db.delete(g)
        elif g:
            g.score = s.score
        else:
            db.add(Grade(student_id=s.student_id, subject=d.subject, term=d.term, score=s.score))
    db.commit()
    return {"saved": len(d.scores)}


def card_data(db: Session, s: Student):
    by: dict = {}
    for g in db.scalars(select(Grade).where(Grade.student_id == s.id)):
        by.setdefault(g.subject, {})[g.term] = g.score
    subjects = []
    for sub in SUBJECTS:
        sc = by.get(sub, {})
        subjects.append(dict(subject=sub, scores={t: sc.get(t) for t in TERMS}, average=round(sum(sc.values()) / len(sc), 2) if sc else None))
    avgs = [x["average"] for x in subjects if x["average"] is not None]
    return dict(subjects=subjects, overall=round(sum(avgs) / len(avgs), 2) if avgs else None)


def pdf_response(buf: BytesIO, name: str):
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{name}"'})


def report_pdf(db: Session, s: Student) -> BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    d, st, buf = card_data(db, s), getSampleStyleSheet(), BytesIO()
    navy = colors.HexColor("#0a2556")
    f = lambda v: "-" if v is None else f"{v:g}"
    rows = [["Subject", *TERMS, "Average"]] + [[x["subject"], *[f(v) for v in x["scores"].values()], f(x["average"])] for x in d["subjects"]]
    rows.append(["Overall average", "", "", "", f(d["overall"])])
    t = Table(rows, repeatRows=1, colWidths=[160, 70, 70, 70, 80])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), .4, colors.lightgrey), ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e8f0fe")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f6f8fc")]),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    info = (f"<b>Name:</b> {s.full_name} &nbsp;&nbsp; <b>No:</b> {s.student_number}<br/>"
            f"<b>Class:</b> {s.school_class.name} ({s.school_class.academic_year}) &nbsp;&nbsp; <b>Born:</b> {s.birth_date:%d/%m/%Y}")
    SimpleDocTemplate(buf, pagesize=A4, topMargin=40, bottomMargin=40, title=f"Report card - {s.full_name}").build([
        Paragraph("HODARI SCHOOL", st["Title"]), Paragraph("Student Report Card", st["Heading2"]), Spacer(1, 8),
        Paragraph(info, st["Normal"]), Spacer(1, 18), t, Spacer(1, 22),
        Paragraph(f"Scores out of 20 - Generated on {date.today():%d/%m/%Y}", st["Italic"])])
    return buf


@router.get("/report-card/{sid}.pdf")
def card(sid: int, db: Session = Depends(get_db)):
    s = db.get(Student, sid)
    if not s:
        raise HTTPException(404, "Student not found")
    return pdf_response(report_pdf(db, s), f"report-card-{s.student_number}.pdf")


def filtered(db: Session, q: str):
    q = q.strip()
    filt = [or_(Student.full_name.ilike(f"%{q}%"), Student.student_number.ilike(f"%{q}%"))] if q else []
    return db.scalars(select(Student).options(joinedload(Student.school_class)).where(*filt).order_by(Student.student_number)).all()


HEAD = ["#", "Student Number", "Full Name", "Class", "Academic Year", "Date of Birth", "Guardian", "Status"]
line = lambda i, s: [i, s.student_number, s.full_name, s.school_class.name, s.school_class.academic_year, f"{s.birth_date:%d/%m/%Y}", s.guardian_name, s.status]


@router.get("/export/students.xlsx")
def export_xlsx(q: str = "", db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(HEAD)
    for i, s in enumerate(filtered(db, q), 1):
        ws.append(line(i, s))
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0A2556"), Alignment(horizontal="center")
    for col, w in zip("ABCDEFGH", [5, 16, 28, 10, 15, 15, 24, 10]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="students.xlsx"'})


@router.get("/export/students.pdf")
def export_pdf(q: str = "", db: Session = Depends(get_db)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    st, buf = getSampleStyleSheet(), BytesIO()
    t = Table([HEAD] + [line(i, s) for i, s in enumerate(filtered(db, q), 1)], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0a2556")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), .3, colors.lightgrey), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f8fc")])]))
    SimpleDocTemplate(buf, pagesize=landscape(A4), title="Students list").build([
        Paragraph("HODARI SCHOOL - Students list", st["Title"]), Paragraph(f"Generated on {date.today():%d/%m/%Y}", st["Italic"]), Spacer(1, 12), t])
    return pdf_response(buf, "students.pdf")
